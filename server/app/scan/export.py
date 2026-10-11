"""Export approved / final grades of one exam to CSV or Excel (.xlsx).

One row per student: a column per question (final points, blank while ungraded), total,
maximum, percent, status. AI suggestions are never exported — only teacher grades and
automatic (bubble / computer) grades. The xlsx has a second sheet "Feedback" with the
approved comment per answer. The number of blank (ungraded) cells is returned in the
`X-FairPlay-Ungraded` header so the UI can warn.
"""
from __future__ import annotations

import asyncio
import csv
import io
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..db import get_session
from ..grading import exam_questions, grade_attempt, num, percent
from ..models import Attempt, Exam, User
from ..security import Principal, require_teacher
from .answer_key import _own_exam

router = APIRouter(prefix="/api", tags=["scan-ai"])


def _safe(v: Any) -> Any:
    """Neutralise spreadsheet formulas in text cells (CSV/Excel injection)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


async def build_rows(session: AsyncSession, e: Exam) -> dict:
    qs = exam_questions(e)
    rows = (await session.execute(
        select(Attempt, User.username, User.name).join(User, User.id == Attempt.student_id)
        .where(Attempt.exam_id == e.id, Attempt.status == "submitted")
        .order_by(User.name, User.username))).all()
    students: list[dict] = []
    feedback: list[dict] = []
    ungraded = 0
    seen: set[int] = set()
    for a, username, name in rows:
        seen.add(a.student_id)
        _, pending, grows = grade_attempt(qs, a.answers, a.grades)
        cells: list[Any] = []
        total = 0.0
        for i, (q, r) in enumerate(zip(qs, grows)):
            g = (a.grades or {}).get(q["id"]) or {}
            if r["pending"]:
                cells.append(None)
                ungraded += 1
            else:
                cells.append(num(float(r["points_awarded"])))
                total += float(r["points_awarded"])
            mis = g.get("misconception") or {}
            feedback.append({
                "student": name, "username": username, "question": i + 1, "prompt": q.get("prompt", "")[:200],
                "points": None if r["pending"] else num(float(r["points_awarded"])), "max": q.get("points", 1),
                "status": "ungraded" if r["pending"] else ("graded" if r["manual"] else "automatic"),
                "feedback": r["comment"] if not r["pending"] else "",
                "misconception": (mis.get("text") if mis.get("tag") == "other" else mis.get("tag")) or "",
                "graded_by": g.get("graded_by") or "", "source": g.get("source") or "",
                "graded_at": g.get("graded_at") or "",
            })
        if a.overall_comment:
            feedback.append({"student": name, "username": username, "question": "overall", "prompt": "",
                             "points": None, "max": None, "status": "", "feedback": a.overall_comment,
                             "misconception": "", "graded_by": "", "source": "", "graded_at": ""})
        students.append({"name": name, "username": username, "cells": cells, "total": num(total),
                         "percent": percent(total, e.total_points), "pending": pending,
                         "status": "final" if pending == 0 else f"incomplete ({pending} ungraded)"})
    missing = 0
    if (getattr(e, "delivery", None) or "") == "paper":
        from ..routes_scan import roster_of  # late import (routes_scan imports app.scan)
        _, roster = await roster_of(session, e)
        for u in roster:
            if u.id not in seen:
                missing += 1
                students.append({"name": u.name, "username": u.username, "cells": [None] * len(qs),
                                 "total": None, "percent": None, "pending": len(qs), "status": "not scanned"})
    return {"questions": qs, "students": students, "feedback": feedback, "ungraded": ungraded, "missing": missing}


def _header(qs: list[dict]) -> list[str]:
    return (["Student", "Username"] + [f"Q{i + 1} ({num(float(q.get('points', 1)))})" for i, q in enumerate(qs)]
            + ["Total", "Max", "Percent", "Status"])


def to_csv(e: Exam, data: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_header(data["questions"]))
    for s in data["students"]:
        w.writerow([_safe(s["name"]), _safe(s["username"])] + ["" if c is None else c for c in s["cells"]]
                   + ["" if s["total"] is None else s["total"], num(float(e.total_points)),
                      "" if s["percent"] is None else s["percent"], s["status"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens UTF-8 names correctly


def to_xlsx(e: Exam, data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    head = _header(data["questions"])
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    blank_fill = PatternFill("solid", fgColor="FDEEE3")
    nq = len(data["questions"])
    for s in data["students"]:
        ws.append([_safe(s["name"]), _safe(s["username"])] + list(s["cells"])
                  + [s["total"], num(float(e.total_points)), s["percent"], s["status"]])
        r = ws.max_row
        for j, v in enumerate(s["cells"]):
            if v is None:
                ws.cell(row=r, column=3 + j).fill = blank_fill
    ws.freeze_panes = "C2"
    widths = [24, 16] + [10] * nq + [9, 7, 9, 22]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    if data["ungraded"] or data["missing"]:
        ws.append([])
        ws.append([f"Warning: {data['ungraded']} answers are not graded yet (blank cells)"
                   + (f"; {data['missing']} students were not scanned" if data["missing"] else "")])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="8A3508")

    fb = wb.create_sheet("Feedback")
    fb.append(["Student", "Username", "Question", "Points", "Max", "Status", "Feedback", "Misconception",
               "Graded by", "Source", "Graded at"])
    for c in fb[1]:
        c.font = Font(bold=True)
    for f in data["feedback"]:
        fb.append([_safe(f["student"]), _safe(f["username"]), f["question"], f["points"], f["max"], f["status"],
                   _safe(f["feedback"]), _safe(f["misconception"]), _safe(f["graded_by"]), f["source"], f["graded_at"]])
    fb.freeze_panes = "C2"
    for i, wdt in enumerate([24, 16, 9, 8, 6, 11, 60, 20, 14, 12, 22], start=1):
        fb.column_dimensions[get_column_letter(i)].width = wdt
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _filename(e: Exam, ext: str) -> str:
    name = "".join(c if c.isalnum() else "-" for c in e.title)[:60].strip("-") or "exam"
    return f"{name}-grades.{ext}"


@router.get("/exams/{exam_id}/grades/export.{fmt}")
async def export_grades(exam_id: str, fmt: str, p: Principal = Depends(require_teacher),
                        session: AsyncSession = Depends(get_session)):
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(404, "Use .csv or .xlsx")
    e = await _own_exam(session, exam_id, p)
    data = await build_rows(session, e)
    if fmt == "csv":
        body, media = to_csv(e, data), "text/csv; charset=utf-8"
    else:
        body = await asyncio.to_thread(to_xlsx, e, data)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(body, media_type=media, headers={
        "Content-Disposition": f'attachment; filename="{_filename(e, fmt)}"', "Cache-Control": "no-store",
        "X-FairPlay-Ungraded": str(data["ungraded"]), "X-FairPlay-Missing": str(data["missing"]),
        "Access-Control-Expose-Headers": "X-FairPlay-Ungraded, X-FairPlay-Missing, Content-Disposition"})
