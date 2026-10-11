"""Phase 2 (d): export approved grades to CSV and Excel."""
import csv
import io

import openpyxl
import pytest

from app import ai

from .test_scan import PAPER_EXAM, _attempt, _scan_class, _wait


@pytest.fixture(autouse=True)
def ai_off(monkeypatch):
    monkeypatch.setenv("FAIRPLAY_AI_PROVIDER", "none")
    monkeypatch.delenv("FAIRPLAY_AI_KEY", raising=False)
    yield
    ai.set_transport(None)


@pytest.fixture
def paper(client, teacher):
    e = client.post("/api/exams", json=dict(PAPER_EXAM, title="=Paper Midterm"), headers=teacher).json()
    client.put(f"/api/exams/{e['id']}/roster", json={"usernames": ["alice", "bob"]}, headers=teacher)
    return e


def test_export_csv_and_xlsx(client, teacher, paper, alice):
    e = paper
    q = [x["id"] for x in e["questions"]]
    # scan alice + bob, then add carol to the roster: she shows as "not scanned"
    _scan_class(client, teacher, e)
    _wait(client, teacher, e["id"])
    client.put(f"/api/exams/{e['id']}/roster", json={"usernames": ["alice", "bob", "carol"]}, headers=teacher)
    al = _attempt(client, teacher, e["id"], "alice")
    for qid, pts, c in [(q[1], 3, "Right"), (q[2], 2, ""), (q[3], 4, "=SUM(A1) good")]:
        r = client.post(f"/api/exams/{e['id']}/scan/answers/{al['attempt_id']}/{qid}/approve",
                        json={"points": pts, "comment": c}, headers=teacher)
        assert r.status_code == 200, r.text

    r = client.get(f"/api/exams/{e['id']}/grades/export.csv", headers=teacher)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert 'filename="Paper-Midterm-grades.csv"' in r.headers["content-disposition"]
    # alice: q5 pending; bob: all 5 pending (bubbles "multiple" and every written box wait for approval)
    ungraded = int(r.headers["x-fairplay-ungraded"])
    assert ungraded == 1 + 5 and r.headers["x-fairplay-missing"] == "1"
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["Student", "Username", "Q1 (2)", "Q2 (3)", "Q3 (2)", "Q4 (5)", "Q5 (4)",
                       "Total", "Max", "Percent", "Status"]
    by = {row[1]: row for row in rows[1:]}
    assert by["alice"][2:7] == ["2", "3", "2", "4", ""]
    assert by["alice"][7:] == ["11", "16", "68.8", "incomplete (1 ungraded)"]
    assert by["bob"][2:7] == ["", "", "", "", ""]
    assert by["carol"][-1] == "not scanned"

    r = client.get(f"/api/exams/{e['id']}/grades/export.xlsx", headers=teacher)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Grades", "Feedback"]
    ws = wb["Grades"]
    assert ws["A1"].value == "Student" and ws["C1"].value == "Q1 (2)" and ws["K1"].value == "Status"
    grid = {ws.cell(row=i, column=2).value: [ws.cell(row=i, column=j).value for j in range(1, 12)]
            for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=2).value}
    assert grid["alice"][2:7] == [2, 3, 2, 4, None]
    assert grid["alice"][7] == 11 and grid["alice"][8] == 16 and grid["alice"][9] == 68.8
    assert grid["carol"][10] == "not scanned"
    warn = [ws.cell(row=i, column=1).value for i in range(1, ws.max_row + 1)]
    assert any(isinstance(v, str) and v.startswith("Warning: 6 answers are not graded") for v in warn)
    fb = wb["Feedback"]
    fb_rows = [[c.value for c in row] for row in fb.iter_rows(min_row=2)]
    a4 = next(r for r in fb_rows if r[1] == "alice" and r[2] == 4)
    assert a4[3] == 4 and a4[5] == "graded" and a4[6] == "'=SUM(A1) good"  # formula neutralised
    assert a4[8] == "teacher" and a4[9] == "manual"
    a5 = next(r for r in fb_rows if r[1] == "alice" and r[2] == 5)
    assert a5[3] is None and a5[5] == "ungraded"
    a1 = next(r for r in fb_rows if r[1] == "alice" and r[2] == 1)
    assert a1[5] == "automatic" and a1[3] == 2

    assert client.get(f"/api/exams/{e['id']}/grades/export.pdf", headers=teacher).status_code == 404
    assert client.get(f"/api/exams/{e['id']}/grades/export.csv", headers=alice).status_code == 403
